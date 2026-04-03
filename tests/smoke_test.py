from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from services.excel_service import ExcelService
from utils.logger import get_logger


def build_bank_file(path: Path):
    rows = [
        ["账户：xxx", None, None, None, None, None],
        ["起止日期：2026-02", None, None, None, None, None],
        ["交易日期", "对方户名", "备注（用途）", "摘要", "借方金额", "贷方金额"],
        ["2026-02-01", "供应商A", "采购付款", "", "1,200.50", ""],
        ["2026-02-03", "供应商B", "", "手续费", "", "(35.2)"],
        [None, None, "合计", None, None, None],
    ]
    df = pd.DataFrame(rows)
    df.to_excel(path, header=False, index=False)


def build_template_file(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "手工日记账#单据头(FBillHead)"
    headers = [
        "单据头(序号)",
        "日记账明细(序号)",
        "*(单据头)业务日期",
        "会计期间",
        "(日记账明细)供应商#名称(Null)",
        "(日记账明细)摘要",
        "(日记账明细)借方金额",
        "(日记账明细)贷方金额",
    ]
    ws.append(headers)
    ws.append([1, 1, "2026-01-01", "2026-01", "期初", "期初", 100, 0])
    wb.save(path)


def main():
    data_dir = Path("tests/tmp")
    data_dir.mkdir(parents=True, exist_ok=True)
    bank = data_dir / "bank.xlsx"
    template = data_dir / "template.xlsx"
    build_bank_file(bank)
    build_template_file(template)

    service = ExcelService(get_logger("test"))
    result = service.process(str(bank), str(template), str(data_dir), print)
    assert Path(result.output_path).exists()

    wb = load_workbook(result.output_path)
    ws = wb["手工日记账#单据头(FBillHead)"]
    assert ws.max_row >= 4
    # 第一条流水借方金额 -> 模板贷方金额
    assert ws.cell(3, 8).value == 1200.5
    # 第二条流水贷方金额 -> 模板借方金额
    assert ws.cell(4, 7).value == -35.2
    assert ws.cell(4, 6).value == "手续费"

    print("smoke test passed:", result.output_path)


if __name__ == "__main__":
    main()
