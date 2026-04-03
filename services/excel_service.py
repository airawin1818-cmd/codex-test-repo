from __future__ import annotations

import copy
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from services.mapping_service import MappingService
from utils.helpers import clean_text, ensure_dir, normalize_text, parse_date, safe_filename

LogFunc = Callable[[str], None]


@dataclass
class ProcessResult:
    output_path: str
    rows_written: int


class ExcelService:
    BANK_HEADER_KEYS = ["交易日期", "对方户名", "备注(用途)", "摘要", "借方金额", "贷方金额"]

    BANK_ALIASES = {
        "交易日期": ["交易日期", "记账日期", "入账日期", "交易时间"],
        "对方户名": ["对方户名", "对方账户名称", "对方名称", "对方单位"],
        "备注(用途)": ["备注(用途)", "备注", "用途", "交易用途", "附言"],
        "摘要": ["摘要", "交易摘要", "说明"],
        "借方金额": ["借方金额", "收入金额", "贷记金额", "借方发生额"],
        "贷方金额": ["贷方金额", "支出金额", "借记金额", "贷方发生额"],
    }

    TEMPLATE_ALIASES = {
        "business_date": ["*(单据头)业务日期", "单据头业务日期", "业务日期"],
        "supplier": ["(日记账明细)供应商#名称(null)", "供应商#名称", "供应商名称"],
        "summary": ["(日记账明细)摘要", "日记账明细摘要", "摘要"],
        "credit": ["(日记账明细)贷方金额", "贷方金额"],
        "debit": ["(日记账明细)借方金额", "借方金额"],
        "head_seq": ["单据头(序号)", "单据头序号", "序号"],
        "detail_seq": ["日记账明细(序号)", "日记账明细序号"],
        "period": ["会计期间", "期间", "单据头会计期间"],
    }

    def __init__(self, logger):
        self.logger = logger

    def process(self, bank_file: str, template_file: str, output_dir: str, log: LogFunc) -> ProcessResult:
        log("开始处理…")
        self.logger.info("处理开始: bank=%s template=%s", bank_file, template_file)

        bank_df = self._load_bank_statement(bank_file, log)
        wb, ws = self._load_template_sheet(template_file, log)

        header_row, template_cols = self._locate_template_columns(ws)
        missing_template = [k for k, v in template_cols.items() if k in ["business_date", "supplier", "summary", "credit", "debit", "head_seq", "detail_seq", "period"] and v is None]
        if missing_template:
            raise ValueError(f"模板缺少必要列: {', '.join(missing_template)}")

        start_row = ws.max_row + 1
        period_format = self._detect_period_format(ws, template_cols["period"], header_row + 1)
        head_seq = self._max_numeric(ws, template_cols["head_seq"], header_row + 1)
        detail_seq = self._max_numeric(ws, template_cols["detail_seq"], header_row + 1)
        log(f"模板起始写入行: {start_row}")

        rows_written = 0
        sample_row_for_style = max(header_row + 1, ws.max_row)

        for _, row in bank_df.iterrows():
            mapped = MappingService.map_record(row.to_dict())
            if not mapped.business_date:
                continue

            target_row = start_row + rows_written
            self._copy_row_style(ws, sample_row_for_style, target_row)

            head_seq += 1
            detail_seq += 1

            ws.cell(target_row, template_cols["head_seq"], head_seq)
            ws.cell(target_row, template_cols["detail_seq"], detail_seq)

            bdate_cell = ws.cell(target_row, template_cols["business_date"], mapped.business_date)
            if ws.cell(sample_row_for_style, template_cols["business_date"]).number_format:
                bdate_cell.number_format = ws.cell(sample_row_for_style, template_cols["business_date"]).number_format

            ws.cell(target_row, template_cols["supplier"], mapped.supplier_name)
            ws.cell(target_row, template_cols["summary"], mapped.summary)
            ws.cell(target_row, template_cols["credit"], mapped.credit_amount)
            ws.cell(target_row, template_cols["debit"], mapped.debit_amount)
            ws.cell(target_row, template_cols["period"], self._format_period(mapped.business_date, period_format))

            rows_written += 1

        if rows_written == 0:
            raise ValueError("未识别到可写入的有效流水数据（交易日期为空或无有效记录）。")

        ensure_dir(output_dir)
        output_name = safe_filename(f"{Path(template_file).stem}_自动填充_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        output_path = str(Path(output_dir) / output_name)
        wb.save(output_path)
        log(f"处理完成，共写入 {rows_written} 行。")

        self.logger.info("处理完成: output=%s rows=%s", output_path, rows_written)
        return ProcessResult(output_path=output_path, rows_written=rows_written)

    def _load_bank_statement(self, bank_file: str, log: LogFunc) -> pd.DataFrame:
        ext = Path(bank_file).suffix.lower()
        if ext not in {".xls", ".xlsx"}:
            raise ValueError("银行流水文件仅支持 .xls / .xlsx")

        engine = "xlrd" if ext == ".xls" else None
        raw = pd.read_excel(bank_file, header=None, dtype=object, engine=engine)
        header_idx, col_map = self._detect_bank_header(raw)

        if header_idx is None:
            raise ValueError("未识别到银行流水表头，请检查文件是否包含交易日期/对方户名等字段。")

        log(f"已识别银行流水表头行: 第 {header_idx + 1} 行")

        headers = [clean_text(v) for v in raw.iloc[header_idx].tolist()]
        data = raw.iloc[header_idx + 1 :].copy()
        data.columns = headers

        renamed: Dict[str, str] = {}
        for std_name, idx in col_map.items():
            source_col = headers[idx]
            renamed[source_col] = std_name

        data = data.rename(columns=renamed)
        for key in self.BANK_HEADER_KEYS:
            if key not in data.columns:
                data[key] = None

        data = data[data["交易日期"].apply(lambda x: parse_date(x) is not None)]
        data = data[~data["交易日期"].astype(str).str.contains("合计|总计|小计", na=False)]

        log(f"已加载银行流水数据: {len(data)} 行有效记录")
        return data.reset_index(drop=True)

    def _detect_bank_header(self, raw_df: pd.DataFrame) -> Tuple[Optional[int], Dict[str, int]]:
        best_score = -1
        best_idx: Optional[int] = None
        best_map: Dict[str, int] = {}

        for i in range(min(len(raw_df), 80)):
            row = [normalize_text(v) for v in raw_df.iloc[i].tolist()]
            current_map: Dict[str, int] = {}
            score = 0
            for std_key, aliases in self.BANK_ALIASES.items():
                alias_norm = [normalize_text(a) for a in aliases]
                for cidx, cell in enumerate(row):
                    if any(a and a in cell for a in alias_norm):
                        current_map[std_key] = cidx
                        score += 1
                        break
            if score > best_score:
                best_score = score
                best_idx = i
                best_map = current_map

        if best_score < 4 or best_idx is None:
            return None, {}
        return best_idx, best_map

    def _load_template_sheet(self, template_file: str, log: LogFunc):
        wb = load_workbook(template_file)
        target_ws = None
        for name in wb.sheetnames:
            n = normalize_text(name)
            if "手工日记账" in n and ("fbillhead" in n or "单据头" in n):
                target_ws = wb[name]
                break
        if target_ws is None:
            raise ValueError("模板页签未找到：请确认存在“手工日记账#单据头(FBillHead)”或近似名称")
        log(f"已识别模板页签: {target_ws.title}")
        return wb, target_ws

    def _locate_template_columns(self, ws: Worksheet) -> Tuple[int, Dict[str, Optional[int]]]:
        best_row = 1
        best_cols: Dict[str, Optional[int]] = {k: None for k in self.TEMPLATE_ALIASES}
        best_score = -1
        for r in range(1, min(ws.max_row, 60) + 1):
            row_values = [normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            cols = {k: None for k in self.TEMPLATE_ALIASES}
            score = 0
            for key, aliases in self.TEMPLATE_ALIASES.items():
                alias_norm = [normalize_text(a) for a in aliases]
                for c, cell in enumerate(row_values, start=1):
                    if any(a and (a == cell or a in cell) for a in alias_norm):
                        cols[key] = c
                        score += 1
                        break
            if score > best_score:
                best_score = score
                best_row = r
                best_cols = cols
        return best_row, best_cols

    def _copy_row_style(self, ws: Worksheet, source_row: int, target_row: int) -> None:
        for col in range(1, ws.max_column + 1):
            source = ws.cell(source_row, col)
            target = ws.cell(target_row, col)
            if source.has_style:
                target._style = copy.copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
            if source.fill and not isinstance(source.fill, PatternFill):
                target.fill = copy.copy(source.fill)

    def _max_numeric(self, ws: Worksheet, col_idx: int, start_row: int) -> int:
        max_val = 0
        for r in range(start_row, ws.max_row + 1):
            v = ws.cell(r, col_idx).value
            try:
                if v is None:
                    continue
                if isinstance(v, str) and re.search(r"\d+", v):
                    num = int(re.search(r"\d+", v).group())
                else:
                    num = int(float(v))
                max_val = max(max_val, num)
            except Exception:
                continue
        return max_val

    def _detect_period_format(self, ws: Worksheet, col_idx: int, start_row: int) -> str:
        for r in range(ws.max_row, start_row - 1, -1):
            text = clean_text(ws.cell(r, col_idx).value)
            if not text:
                continue
            if re.match(r"^\d{4}-\d{1,2}$", text):
                return "yyyy-mm"
            if re.match(r"^\d{6}$", text):
                return "yyyymm"
            if re.match(r"^\d{4}年\d{1,2}月$", text):
                return "cn"
        return "yyyy-mm"

    def _format_period(self, dt: datetime, fmt: str) -> str:
        if fmt == "yyyymm":
            return dt.strftime("%Y%m")
        if fmt == "cn":
            return f"{dt.year}年{dt.month}月"
        return dt.strftime("%Y-%m")
